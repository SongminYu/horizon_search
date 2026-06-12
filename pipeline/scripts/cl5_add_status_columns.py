"""Retrofit CL5_Topics_and_Projects.xlsx with the same explicit status columns
as the CL6 workbook, so both files can be filtered by data regime later.

Edits IN PLACE (appends columns at the end, fills empty Keywords cells);
does NOT regenerate rows, so any manual edits in existing cells survive.

Appended to the Topics sheet:
  - Call Status                   : Open / Closed / Forthcoming (from Opening Date / Deadline vs today)
  - Award Status                  : Signed / Not yet signed     (has CORDIS projects?)
  - Funded Projects (signed)
  - Signed EU Contribution (EUR M)
Also fills the (previously empty) Keywords column from cached topicDetails JSON.
"""
import datetime
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"
CACHE = ROOT / "raw" / "topic_details"
XLSX = ROOT.parent / "data" / "CL5_Topics_and_Projects.xlsx"
TODAY = datetime.date.today().isoformat()

proj = pd.read_pickle(PARSED / "cl5_projects.pkl")

def to_float(x):
    if pd.isna(x):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None

proj["_ec"] = proj["ecMaxContribution"].apply(to_float)
proj_by_topic = proj.groupby("topics").agg(funded=("id", "count"), sum_ec=("_ec", "sum")).to_dict("index")

def get_keywords(tid):
    p = CACHE / f"{tid}.json"
    if not p.exists():
        return None
    try:
        td = json.loads(p.read_text()).get("TopicDetails", {})
    except json.JSONDecodeError:
        return None
    kw = td.get("keywords") or []
    return "; ".join(kw) if kw else None

def call_status(opening, deadline):
    if not opening and not deadline:
        return None
    if opening and opening > TODAY:
        return "Forthcoming"
    if deadline and deadline < TODAY:
        return "Closed"
    if opening and opening <= TODAY:
        return "Open"
    return None

wb = load_workbook(XLSX)
ws = wb["Topics"]
headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}

NEW_COLS = ["Call Status", "Award Status", "Funded Projects (signed)", "Signed EU Contribution (EUR M)"]
already = [h for h in NEW_COLS if h in headers]
if already:
    print(f"Columns already present, will overwrite their values: {already}")

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="2C5D99")
for h in NEW_COLS:
    if h not in headers:
        c = ws.max_column + 1
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 14
        headers[h] = c

col_tid = headers["Topic ID"]
col_open = headers["Opening Date"]
col_dl = headers["Deadline"]
col_kw = headers["Keywords"]

n_kw_filled = 0
for r in range(2, ws.max_row + 1):
    tid = ws.cell(row=r, column=col_tid).value
    if not tid:
        continue
    p = proj_by_topic.get(tid, {})
    funded = int(p.get("funded", 0) or 0)
    opening = ws.cell(row=r, column=col_open).value
    deadline = ws.cell(row=r, column=col_dl).value
    opening = str(opening)[:10] if opening else None
    deadline = str(deadline)[:10] if deadline else None

    ws.cell(row=r, column=headers["Call Status"], value=call_status(opening, deadline))
    ws.cell(row=r, column=headers["Award Status"], value="Signed" if funded else "Not yet signed")
    ws.cell(row=r, column=headers["Funded Projects (signed)"], value=funded)
    ws.cell(row=r, column=headers["Signed EU Contribution (EUR M)"],
            value=round((p.get("sum_ec") or 0) / 1e6, 3) if funded else 0)

    if ws.cell(row=r, column=col_kw).value in (None, ""):
        kw = get_keywords(tid)
        if kw:
            ws.cell(row=r, column=col_kw, value=kw)
            n_kw_filled += 1

wb.save(XLSX)
print(f"Saved {XLSX}")
print(f"Rows processed: {ws.max_row - 1}; Keywords filled: {n_kw_filled}")
