"""Append/refresh a "Participants" sheet (one row = one organisation in one project,
with that organisation's own EU money) in a Topics+Projects workbook.

Usage: python3 pipeline/scripts/add_participants_sheet.py <projects_pkl_stem> <workbook>
  e.g. python3 pipeline/scripts/add_participants_sheet.py cl5_projects CL5_Topics_and_Projects.xlsx
       python3 pipeline/scripts/add_participants_sheet.py cl6_projects CL6_Topics_and_Projects.xlsx
       python3 pipeline/scripts/add_participants_sheet.py miss_projects MISS_Topics_and_Projects.xlsx

Source: CORDIS organization.xlsx
  - ecContribution    = EU contribution allocated to this partner in this project
  - netEcContribution = same net of financial-support-to-third-parties flows
  - activityType: HES higher education / REC research org / PRC private company /
                  PUB public body / OTH other
Only the Participants sheet is touched; Topics/Projects sheets (incl. manual edits) stay as-is.
"""
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"

pkl_stem, wb_name = sys.argv[1], sys.argv[2]
RAW_SUBDIR = sys.argv[3] if len(sys.argv) > 3 else "cordis_he"
proj = pd.read_pickle(PARSED / f"{pkl_stem}.pkl")
XLSX = ROOT.parent / "data" / wb_name

proj_meta = proj.set_index("id")[["topics", "acronym"]].to_dict("index")
pids = set(proj_meta)

org = pd.read_excel(ROOT / "raw" / RAW_SUBDIR / "organization.xlsx")
org = org[org["projectID"].isin(pids)].copy()

def to_float(x):
    if pd.isna(x):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None

for col in ["ecContribution", "netEcContribution", "totalCost"]:
    org[col] = org[col].apply(to_float)
org["order"] = pd.to_numeric(org["order"], errors="coerce")
org = org.sort_values(["projectID", "order"])

rows = []
for _, r in org.iterrows():
    pid = r["projectID"]
    m = proj_meta[pid]
    rows.append({
        "Topic ID": m["topics"],
        "Project Acronym": m["acronym"],
        "Grant Agreement ID": pid,
        "Organisation Name": r["name"],
        "Short Name": r["shortName"],
        "Organisation ID": r["organisationID"],
        "Role": r["role"],
        "Country": r["country"],
        "City": r["city"],
        "Activity Type": r["activityType"],
        "SME": r["SME"],
        "EC Contribution (EUR)": r["ecContribution"],
        "Net EC Contribution (EUR)": r["netEcContribution"],
        "Total Cost (EUR)": r["totalCost"],
        "End of Participation": r["endOfParticipation"],
    })
df = pd.DataFrame(rows).sort_values(["Topic ID", "Grant Agreement ID"], kind="stable").reset_index(drop=True)
print(f"Participants rows: {len(df)} across {df['Grant Agreement ID'].nunique()} projects")
print(f"EC Contribution sum: EUR {df['EC Contribution (EUR)'].sum()/1e9:.2f}B")

WIDTHS = {
    "Topic ID": 36, "Project Acronym": 16, "Grant Agreement ID": 14, "Organisation Name": 50,
    "Short Name": 16, "Organisation ID": 14, "Role": 16, "Country": 9, "City": 18,
    "Activity Type": 10, "SME": 8, "EC Contribution (EUR)": 15, "Net EC Contribution (EUR)": 15,
    "Total Cost (EUR)": 15, "End of Participation": 12,
}

wb = load_workbook(XLSX)
if "Participants" in wb.sheetnames:
    del wb["Participants"]
ws = wb.create_sheet("Participants")
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="B66A1F")
for c_idx, h in enumerate(df.columns, start=1):
    cell = ws.cell(row=1, column=c_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = WIDTHS.get(h, 14)
for r_idx, rec in enumerate(df.to_dict("records"), start=2):
    for c_idx, h in enumerate(df.columns, start=1):
        v = rec.get(h)
        if v is not None and not isinstance(v, (list, dict)) and pd.isna(v):
            v = None
        ws.cell(row=r_idx, column=c_idx, value=v)
ws.freeze_panes = "A2"
wb.save(XLSX)
print(f"Saved {XLSX} (Participants sheet: {len(df)} rows x {len(df.columns)} cols)")
