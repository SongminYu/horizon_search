"""Merge CORDIS + F&T Portal data into the final Excel workbook.

Output: CL5_Topics_and_Projects.xlsx
        with Sheet 1 = Topics (rows = unique CL5 topics 2021-2026)
             Sheet 2 = Projects (rows = funded projects, one per row)
"""
import json
import re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"
OUT = ROOT.parent / "data" / "CL5_Topics_and_Projects.xlsx"

# ---------- Load all sources ----------
cordis_topics = pd.read_pickle(PARSED / "topics_from_cordis.pkl")          # 390 rows
ft_main       = pd.read_pickle(PARSED / "topics_ft_details.pkl")           # 390 rows
ft_2025       = pd.read_pickle(PARSED / "topics_ft_details_2025_2026.pkl") # 76 rows
ft_2026       = pd.read_pickle(PARSED / "topics_ft_details_2026wp.pkl")    # 42 rows
proj          = pd.read_pickle(PARSED / "cl5_projects.pkl")                # 932 rows

# Combine F&T details
ft_all = pd.concat([ft_main, ft_2025, ft_2026], ignore_index=True)
ft_all = ft_all.drop_duplicates(subset=["topic_id"], keep="first")
print(f"Unique topics with F&T details: {len(ft_all)}")

# ---------- Build Sheet 1 (Topics) ----------
DEST_MAP = {
    "D1": "Climate sciences and responses",
    "D2": "Cross-sectoral solutions for the climate transition",
    "D3": "Sustainable, secure and competitive energy supply",
    "D4": "Efficient, sustainable and inclusive energy use",
    "D5": "Clean and competitive solutions for all transport modes",
    "D6": "Safe, resilient transport and smart mobility services for passengers and goods",
}

# Topic ID parser (same patterns as earlier)
pat_old = re.compile(r"^HORIZON-CL5-(\d{4})-(D\d)-(\d+)-(\d+)(?:-(.+))?$")
pat_new = re.compile(r"^HORIZON-CL5-(\d{4})-(\d+)-(D\d)-(\d+)(?:-(.+))?$")
pat_iba = re.compile(r"^HORIZON-CL5-(\d{4})-([A-Z0-9]+)-IBA(?:-(\d+))?$")

def parse_id(t):
    m = pat_old.match(t)
    if m:
        return m.group(1), m.group(2), m.group(3)
    m = pat_new.match(t)
    if m:
        return m.group(1), m.group(3), m.group(2)
    m = pat_iba.match(t)
    if m:
        return m.group(1), None, m.group(2)
    return None, None, None

# Start from full set of topic IDs (union of CORDIS and FT-only)
all_topic_ids = sorted(set(cordis_topics["topic_id"]) | set(ft_all["topic_id"]))
print(f"Total unique topics: {len(all_topic_ids)}")

# project aggregates by topic
def to_float(x):
    if pd.isna(x): return None
    try: return float(str(x).replace(",", "."))
    except ValueError: return None
proj["_ec"] = proj["ecMaxContribution"].apply(to_float)

proj_by_topic = proj.groupby("topics").agg(
    funded=("id", "count"),
    sum_ec=("_ec", "sum"),
).to_dict("index")

cordis_t_map = cordis_topics.set_index("topic_id").to_dict("index")
ft_map = ft_all.set_index("topic_id").to_dict("index")

PORTAL_TPL = "https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/opportunities/topic-details/{tid}"

rows = []
for tid in all_topic_ids:
    year, dest, call_num = parse_id(tid)
    if year not in {"2021", "2022", "2023", "2024", "2025", "2026"}:
        continue  # filter to 2021-2026
    c = cordis_t_map.get(tid, {})
    f = ft_map.get(tid, {})
    p = proj_by_topic.get(tid, {})

    # Prefer F&T fields; fall back to CORDIS
    title = f.get("ft_title") or c.get("topic_title")
    call_id = f.get("ft_call_id") or c.get("master_call")
    call_title = f.get("ft_call_title")
    topic_type = f.get("ft_type") or c.get("topic_type")
    if pat_iba.match(tid):
        topic_type = "IBA"

    rows.append({
        "Framework": "Horizon Europe",
        "Year": year,
        "Cluster/Pillar": "CL5",
        "Destination": dest,
        "Destination Name": DEST_MAP.get(dest) if dest else None,
        "Call ID": call_id,
        "Call Title": call_title,
        "Topic ID": tid,
        "Topic Title": title,
        "Topic Type": topic_type,
        "TRL": f.get("trl"),
        "EU Contribution per Project (EUR M)": f.get("budget_per_project_M"),
        "Indicative Budget (EUR M)": f.get("budget_total_M"),
        "Expected Number of Projects": f.get("expected_grants"),
        "Opening Date": f.get("opening_date"),
        "Deadline": f.get("deadline_date"),
        "Stage": f.get("stage"),
        "Keywords": None,
        "EU Portal Link": PORTAL_TPL.format(tid=tid.lower()),
        "Work Programme PDF": None,
        "Notes": (
            f"Funded projects: {p.get('funded',0)}; "
            f"Total EU contribution: EUR {p.get('sum_ec',0)/1e6:.2f}M"
            if p else "No funded projects yet"
        ),
    })

topics_df = pd.DataFrame(rows)
# Year ordering keeps 2021..2026 ascending, destination ordering D1..D6
topics_df = topics_df.sort_values(["Year", "Destination", "Topic ID"], na_position="last").reset_index(drop=True)
print(f"Topics rows: {len(topics_df)}")
print(f"By year:\n{topics_df['Year'].value_counts().sort_index()}")

# ---------- Build Sheet 2 (Projects) ----------
PROJECT_PORTAL = "https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/how-to-participate/projects-results/project-details/{rcn}"
CORDIS_LINK = "https://cordis.europa.eu/project/id/{id}"

# weblink lookup for project homepages
wl = pd.read_excel(ROOT / "raw" / "cordis_he" / "webLink.xlsx")
homepage = wl[(wl["type"] == "relatedWebsite") & (wl["represents"] == "project")] \
    .groupby("projectID")["physUrl"].first().to_dict()
org = pd.read_excel(ROOT / "raw" / "cordis_he" / "organization.xlsx", usecols=["projectID", "name", "country", "role"])
coords = org[org["role"] == "coordinator"].drop_duplicates("projectID").set_index("projectID")
counts = org.groupby("projectID").size().to_dict()
countries = org.groupby("projectID")["country"].apply(lambda s: ",".join(sorted(set(s.dropna())))).to_dict()

def fmt_eu_amount(x):
    v = to_float(x)
    return v

proj_rows = []
for _, r in proj.iterrows():
    pid = r["id"]
    proj_rows.append({
        "Topic ID": r["topics"],
        "Project Acronym": r["acronym"],
        "Project Full Name": r["title"],
        "Grant Agreement ID": pid,
        "Coordinator": coords.loc[pid]["name"] if pid in coords.index else None,
        "Coordinator Country": coords.loc[pid]["country"] if pid in coords.index else None,
        "Number of Participants": counts.get(pid),
        "Participant Countries": countries.get(pid),
        "Start Date": r["startDate"],
        "End Date": r["endDate"],
        "EU Contribution (EUR)": fmt_eu_amount(r["ecMaxContribution"]),
        "Total Cost (EUR)": fmt_eu_amount(r["totalCost"]),
        "Status": r["status"],
        "Project Website": homepage.get(pid),
        "CORDIS Link": CORDIS_LINK.format(id=pid),
        "EU Portal Link": PROJECT_PORTAL.format(rcn=r["rcn"]),
        "Objective": r["objective"],
        "Notes": None,
    })

projects_df = pd.DataFrame(proj_rows)
projects_df = projects_df.sort_values(["Topic ID", "Project Acronym"]).reset_index(drop=True)
print(f"Project rows: {len(projects_df)}")

# ---------- Write to existing workbook (preserve headers/formatting) ----------
wb = load_workbook(OUT)

def write_sheet(ws, df):
    # clear existing data rows (keep row 1 headers)
    max_existing = ws.max_row
    if max_existing > 1:
        ws.delete_rows(2, max_existing - 1)
    # ensure header order matches
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    missing = [h for h in headers if h not in df.columns]
    if missing:
        print(f"  ! Missing columns in df: {missing}")
    # write
    for r_idx, rec in enumerate(df.to_dict("records"), start=2):
        for c_idx, h in enumerate(headers, start=1):
            v = rec.get(h)
            if pd.isna(v) if not isinstance(v, (list, dict)) else False:
                v = None
            ws.cell(row=r_idx, column=c_idx, value=v)

write_sheet(wb["Topics"], topics_df)
write_sheet(wb["Projects"], projects_df)

wb.save(OUT)
print(f"\nSaved: {OUT}")
print(f"  Topics sheet: {len(topics_df)} rows")
print(f"  Projects sheet: {len(projects_df)} rows")
