"""Build a Topics-only workbook for programmes whose awarded-project data has
no open bulk source (LIFE, Innovation Fund): same 25-column Topics sheet as the
Horizon workbooks, but no Projects/Participants/Organisations sheets.

Usage: python3 pipeline/scripts/topics_only_merge.py LIFE
       python3 pipeline/scripts/topics_only_merge.py INNOVFUND

Award-status columns are left blank (project layer not public: CINEA Qlik
dashboards only); Notes column records this.
"""
import datetime
import json
import re
import sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"
CACHE = ROOT / "raw" / "topic_details"
TODAY = datetime.date.today().isoformat()

PROG = sys.argv[1]
CONF = {
    "LIFE": {
        "pkl": "life_topics_ft_details.pkl",
        "out": "LIFE_Topics.xlsx",
        "framework": "LIFE 2021-2027",
        "dest_names": {
            "NAT": "Nature and Biodiversity sub-programme",
            "ENV": "Circular Economy and Quality of Life sub-programme",
            "CLIMA": "Climate Change Mitigation and Adaptation sub-programme",
            "CET": "Clean Energy Transition sub-programme",
            "OTHER": "Horizontal (preparatory / technical assistance / NGO operating grants)",
        },
        "note": "Project-level data not public in bulk (CINEA LIFE dashboard only; LIFE is not in CORDIS)",
    },
    "INNOVFUND": {
        "pkl": "innovfund_topics_ft_details.pkl",
        "out": "INNOVFUND_Topics.xlsx",
        "framework": "Innovation Fund (EU ETS revenues)",
        "dest_names": {
            "LSC": "Large-scale projects call",
            "SSC": "Small-scale projects call",
            "NZT": "Net-zero technologies call",
            "BATTERIES": "Batteries call",
            "AUCTION": "Hydrogen Bank auction",
            "OTHER": "Other calls / support actions",
        },
        "note": "Project-level data not public in bulk (CINEA Innovation Fund dashboard only)",
    },
}[PROG]

ft = pd.read_pickle(PARSED / CONF["pkl"]).drop_duplicates(subset=["topic_id"])
ft = ft[ft["found"]]
print(f"{PROG}: {len(ft)} topics with details")

def parse_dest(tid):
    body = re.sub(rf"^{PROG}-", "", re.sub(r"20\d\d-?", "", tid)).strip("-")
    toks = body.split("-")
    for t in toks[:3]:
        u = t.upper()
        if u in CONF["dest_names"]:
            return u
        if PROG == "LIFE" and u in {"NAT", "ENV", "CLIMA", "CET"}:
            return u
        if PROG == "INNOVFUND" and ("AUC" in u):
            return "AUCTION"
    return "OTHER"

def get_keywords(tid):
    p = CACHE / f"{tid}.json"
    if not p.exists():
        return None
    try:
        td = json.loads(p.read_text()).get("TopicDetails", {})
    except json.JSONDecodeError:
        return None
    kw = td.get("keywords") or []
    return "; ".join(kw) if kw else None

def call_status(opening, deadline):
    opening = None if (opening is None or pd.isna(opening)) else str(opening)[:10]
    deadline = None if (deadline is None or pd.isna(deadline)) else str(deadline)[:10]
    if not opening and not deadline:
        return None
    if opening and opening > TODAY:
        return "Forthcoming"
    if deadline and deadline < TODAY:
        return "Closed"
    if opening and opening <= TODAY:
        return "Open"
    return None

PORTAL_TPL = "https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/opportunities/topic-details/{tid}"

rows = []
for _, f in ft.iterrows():
    tid = f["topic_id"]
    m = re.search(r"(20\d\d)", tid)
    year = m.group(1) if m else None
    if year not in {"2021", "2022", "2023", "2024", "2025", "2026"}:
        continue
    dest = parse_dest(tid)
    rows.append({
        "Framework": CONF["framework"],
        "Year": year,
        "Cluster/Pillar": PROG,
        "Destination": dest,
        "Destination Name": CONF["dest_names"].get(dest),
        "Call ID": f.get("ft_call_id"),
        "Call Title": f.get("ft_call_title"),
        "Topic ID": tid,
        "Topic Title": f.get("ft_title"),
        "Topic Type": f.get("ft_type"),
        "TRL": f.get("trl"),
        "EU Contribution per Project (EUR M)": f.get("budget_per_project_M"),
        "Indicative Budget (EUR M)": f.get("budget_total_M"),
        "Expected Number of Projects": f.get("expected_grants"),
        "Opening Date": f.get("opening_date"),
        "Deadline": f.get("deadline_date"),
        "Stage": f.get("stage"),
        "Keywords": get_keywords(tid),
        "EU Portal Link": PORTAL_TPL.format(tid=tid.lower()),
        "Work Programme PDF": None,
        "Notes": CONF["note"],
        "Call Status": call_status(f.get("opening_date"), f.get("deadline_date")),
        "Award Status": None,
        "Funded Projects (signed)": None,
        "Signed EU Contribution (EUR M)": None,
    })

df = pd.DataFrame(rows).sort_values(["Year", "Destination", "Topic ID"]).reset_index(drop=True)
print(f"Topics rows: {len(df)}")
print(df.groupby(["Year", "Destination"]).size().unstack(fill_value=0).to_string())

WIDTHS = {
    "Framework": 22, "Year": 8, "Cluster/Pillar": 12, "Destination": 12, "Destination Name": 44,
    "Call ID": 30, "Call Title": 40, "Topic ID": 36, "Topic Title": 60, "Topic Type": 12,
    "TRL": 10, "EU Contribution per Project (EUR M)": 16, "Indicative Budget (EUR M)": 14,
    "Expected Number of Projects": 12, "Opening Date": 13, "Deadline": 13, "Stage": 14,
    "Keywords": 24, "EU Portal Link": 50, "Work Programme PDF": 24, "Notes": 50,
    "Call Status": 14, "Award Status": 14, "Funded Projects (signed)": 12,
    "Signed EU Contribution (EUR M)": 14,
}

wb = Workbook()
ws = wb.active
ws.title = "Topics"
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="2C5D99")
for c_idx, h in enumerate(df.columns, start=1):
    cell = ws.cell(row=1, column=c_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = WIDTHS.get(h, 16)
for r_idx, rec in enumerate(df.to_dict("records"), start=2):
    for c_idx, h in enumerate(df.columns, start=1):
        v = rec.get(h)
        if v is not None and not isinstance(v, (list, dict)) and pd.isna(v):
            v = None
        ws.cell(row=r_idx, column=c_idx, value=v)
ws.freeze_panes = "A2"
out = ROOT.parent / "data" / CONF["out"]
wb.save(out)
print(f"Saved {out}: {len(df)} rows x {len(df.columns)} cols")
