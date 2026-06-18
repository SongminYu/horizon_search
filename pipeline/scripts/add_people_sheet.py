"""Append/refresh a "People" sheet: one row per researcher linked (via OpenAlex
authorship on grant-funded works) to this workbook's projects, with their ORCID,
main institution, output counts, topics, and the projects they appear in.

Usage: python3 pipeline/scripts/add_people_sheet.py <workbook>
  e.g. python3 pipeline/scripts/add_people_sheet.py CL5_Topics_and_Projects.xlsx

Source: parsed/<stem>_people.pkl (run fetch_people.py first).
Sorted by Lead Works (first/last/corresponding-author output) then Works, so the
core PIs sit on top and mega-paper co-authors sink. Only the People sheet is touched.
"""
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"

wb_name = sys.argv[1]
stem = wb_name.split("_Topics")[0].lower()
XLSX = ROOT.parent / "data" / wb_name
df = pd.read_pickle(PARSED / f"{stem}_people.pkl")

print(f"People: {len(df):,} researchers | ORCID {df['ORCID'].notna().mean():.0%} | "
      f"in >=2 projects {int((df['Projects']>=2).sum()):,}")
print("Top 5 by lead output:")
print(df.head(5)[["Name", "Lead Works", "Works", "Projects", "Main Institution"]].to_string(index=False))

WIDTHS = {
    "Author ID": 13, "Name": 26, "ORCID": 20, "Country": 8, "Main Institution": 42,
    "Works": 8, "Lead Works": 9, "Projects": 9, "Lead Projects": 9, "Top Topics": 50,
    "Project Acronyms": 60, "Project Grant IDs": 40,
}

wb = load_workbook(XLSX)
if "People" in wb.sheetnames:
    del wb["People"]
ws = wb.create_sheet("People")
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F7A7A")
for c_idx, h in enumerate(df.columns, start=1):
    cell = ws.cell(row=1, column=c_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = WIDTHS.get(h, 14)
for r_idx, rec in enumerate(df.to_dict("records"), start=2):
    for c_idx, h in enumerate(df.columns, start=1):
        v = rec.get(h)
        if v is not None and not isinstance(v, (list, dict)) and pd.isna(v):
            v = None
        ws.cell(row=r_idx, column=c_idx, value=v)
ws.freeze_panes = "A2"
wb.save(XLSX)
print(f"Saved {XLSX} (People sheet: {len(df)} rows x {len(df.columns)} cols)")
