"""Merge EU Missions CORDIS + F&T Portal data into MISS_Topics_and_Projects.xlsx.

Same 25/18-column layout as the CL5/CL6 workbooks. Semantics differences:
  - Cluster/Pillar = "MISS"
  - Destination      = mission code (CLIMA / OCEAN / CIT / SOIL / CANCER / NEB / ...,
                       compound codes like CLIMA-OCEAN-SOIL for joint topics)
  - Destination Name = mission full name (joint topics get joined short names)
All missions included; filter Destination to CLIMA/OCEAN/CIT/SOIL for the
environment-only view.
"""
import datetime
import json
import re
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"
CACHE = ROOT / "raw" / "topic_details"
OUT = ROOT.parent / "data" / "MISS_Topics_and_Projects.xlsx"
TODAY = datetime.date.today().isoformat()

cordis_topics = pd.read_pickle(PARSED / "miss_topics_from_cordis.pkl")
ft_all = pd.read_pickle(PARSED / "miss_topics_ft_details.pkl")
proj = pd.read_pickle(PARSED / "miss_projects.pkl")

ft_all = ft_all.drop_duplicates(subset=["topic_id"], keep="first")
ft_found_ids = set(ft_all.loc[ft_all["found"], "topic_id"])
print(f"F&T details rows: {len(ft_all)} (found: {len(ft_found_ids)})")

MISSION_NAME = {
    "CLIMA": "Adaptation to Climate Change",
    "OCEAN": "Restore our Ocean and Waters by 2030",
    "CIT": "Climate-Neutral and Smart Cities",
    "CITIES": "Climate-Neutral and Smart Cities",
    "SOIL": "A Soil Deal for Europe",
    "CANCER": "Mission Cancer",
    "UNCAN": "Mission Cancer (UNCAN.eu)",
    "NEB": "New European Bauhaus Facility",
    "CROSS": "Cross-mission activities",
    "COOR": "Mission coordination and support",
    "NCP": "Mission National Contact Points",
}
MISSION_SHORT = {
    "CLIMA": "Climate Adaptation", "OCEAN": "Ocean & Waters", "CIT": "Climate-Neutral Cities",
    "CITIES": "Climate-Neutral Cities", "SOIL": "Soil Deal", "CANCER": "Cancer",
    "NEB": "New European Bauhaus",
}

def dest_name(code):
    if not code:
        return None
    if code in MISSION_NAME:
        return MISSION_NAME[code]
    parts = code.split("-")
    if all(p in MISSION_SHORT for p in parts):
        return "Joint: " + " + ".join(MISSION_SHORT[p] for p in parts)
    return None

MISS_SEG = r"[A-Z]+(?:-[A-Z]+)*"
pat_old = re.compile(rf"^HORIZON-MISS-(\d{{4}})-({MISS_SEG}?)-(\d+)-(\d+)(?:-(.+))?$")
pat_new = re.compile(rf"^HORIZON-MISS-(\d{{4}})-(\d+)-({MISS_SEG}?)-(\d+)(?:-(.+))?$")
pat_iba = re.compile(rf"^HORIZON-MISS-(\d{{4}})-(?:(\d+)-)?({MISS_SEG}?)-IBA(?:-(\d+))?$")
pat_sga = re.compile(rf"^HORIZON-MISS-(\d{{4}})-(?:(\d+)-)?({MISS_SEG}?)-SGA(?:-(\d+))?$")

def parse_id(t):
    """-> (year, mission_code, kind)"""
    m = pat_iba.match(t)
    if m:
        return m.group(1), m.group(3), "IBA"
    m = pat_sga.match(t)
    if m:
        return m.group(1), m.group(3), "SGA"
    m = pat_old.match(t)
    if m:
        return m.group(1), m.group(2), "standard"
    m = pat_new.match(t)
    if m:
        return m.group(1), m.group(3), "standard"
    return None, None, None

all_topic_ids = set(cordis_topics["topic_id"]) | ft_found_ids
all_topic_ids -= {t[:-len("-two-stage")] for t in all_topic_ids if t.endswith("-two-stage")
                  and t[:-len("-two-stage")] not in set(cordis_topics["topic_id"])}
all_topic_ids = sorted(all_topic_ids)
print(f"Total unique topic IDs: {len(all_topic_ids)}")

def get_keywords(tid):
    p = CACHE / f"{tid}.json"
    if not p.exists():
        return None
    try:
        td = json.loads(p.read_text()).get("TopicDetails", {})
    except json.JSONDecodeError:
        return None
    kw = td.get("keywords") or []
    return "; ".join(kw) if kw else None

def call_status(opening, deadline):
    opening = None if (opening is None or pd.isna(opening)) else str(opening)[:10]
    deadline = None if (deadline is None or pd.isna(deadline)) else str(deadline)[:10]
    if not opening and not deadline:
        return None
    if opening and opening > TODAY:
        return "Forthcoming"
    if deadline and deadline < TODAY:
        return "Closed"
    if opening and opening <= TODAY:
        return "Open"
    return None

def to_float(x):
    if pd.isna(x):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None

proj["_ec"] = proj["ecMaxContribution"].apply(to_float)
proj_by_topic = proj.groupby("topics").agg(funded=("id", "count"), sum_ec=("_ec", "sum")).to_dict("index")
cordis_t_map = cordis_topics.set_index("topic_id").to_dict("index")
ft_map = ft_all.set_index("topic_id").to_dict("index")

PORTAL_TPL = "https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/opportunities/topic-details/{tid}"

rows = []
for tid in all_topic_ids:
    year, dest, kind = parse_id(tid)
    if year not in {"2021", "2022", "2023", "2024", "2025", "2026"}:
        continue
    c = cordis_t_map.get(tid, {})
    f = ft_map.get(tid, {})
    p = proj_by_topic.get(tid, {})
    title = f.get("ft_title") or c.get("topic_title")
    call_id = f.get("ft_call_id") or c.get("master_call")
    topic_type = f.get("ft_type") or c.get("topic_type")
    if kind in ("IBA", "SGA"):
        topic_type = kind
    funded = p.get("funded", 0)
    rows.append({
        "Framework": "Horizon Europe",
        "Year": year,
        "Cluster/Pillar": "MISS",
        "Destination": dest,
        "Destination Name": dest_name(dest),
        "Call ID": call_id,
        "Call Title": f.get("ft_call_title"),
        "Topic ID": tid,
        "Topic Title": title,
        "Topic Type": topic_type,
        "TRL": f.get("trl"),
        "EU Contribution per Project (EUR M)": f.get("budget_per_project_M"),
        "Indicative Budget (EUR M)": f.get("budget_total_M"),
        "Expected Number of Projects": f.get("expected_grants"),
        "Opening Date": f.get("opening_date"),
        "Deadline": f.get("deadline_date"),
        "Stage": f.get("stage"),
        "Keywords": get_keywords(tid),
        "EU Portal Link": PORTAL_TPL.format(tid=tid.lower()),
        "Work Programme PDF": None,
        "Notes": None,
        "Call Status": call_status(f.get("opening_date"), f.get("deadline_date")),
        "Award Status": "Signed" if funded else "Not yet signed",
        "Funded Projects (signed)": funded if funded else 0,
        "Signed EU Contribution (EUR M)": round(p.get("sum_ec", 0) / 1e6, 3) if funded else 0,
    })

topics_df = pd.DataFrame(rows).sort_values(["Year", "Destination", "Topic ID"], na_position="last").reset_index(drop=True)
print(f"Topics rows: {len(topics_df)}")
print(pd.crosstab(topics_df["Year"], topics_df["Award Status"]))

PROJECT_PORTAL = "https://ec.europa.eu/info/funding-tenders/opportunities/portal/screen/how-to-participate/projects-results/project-details/{rcn}"
CORDIS_LINK = "https://cordis.europa.eu/project/id/{id}"

wl = pd.read_excel(ROOT / "raw" / "cordis_he" / "webLink.xlsx")
homepage = (wl[(wl["type"] == "relatedWebsite") & (wl["represents"] == "project")]
            .groupby("projectID")["physUrl"].first().to_dict())
org = pd.read_excel(ROOT / "raw" / "cordis_he" / "organization.xlsx", usecols=["projectID", "name", "country", "role"])
coords = org[org["role"] == "coordinator"].drop_duplicates("projectID").set_index("projectID")
counts = org.groupby("projectID").size().to_dict()
countries = org.groupby("projectID")["country"].apply(lambda s: ",".join(sorted(set(s.dropna())))).to_dict()

proj_rows = []
for _, r in proj.iterrows():
    pid = r["id"]
    proj_rows.append({
        "Topic ID": r["topics"],
        "Project Acronym": r["acronym"],
        "Project Full Name": r["title"],
        "Grant Agreement ID": pid,
        "Coordinator": coords.loc[pid]["name"] if pid in coords.index else None,
        "Coordinator Country": coords.loc[pid]["country"] if pid in coords.index else None,
        "Number of Participants": counts.get(pid),
        "Participant Countries": countries.get(pid),
        "Start Date": r["startDate"],
        "End Date": r["endDate"],
        "EU Contribution (EUR)": to_float(r["ecMaxContribution"]),
        "Total Cost (EUR)": to_float(r["totalCost"]),
        "Status": r["status"],
        "Project Website": homepage.get(pid),
        "CORDIS Link": CORDIS_LINK.format(id=pid),
        "EU Portal Link": PROJECT_PORTAL.format(rcn=r["rcn"]),
        "Objective": r["objective"],
        "Notes": None,
    })
projects_df = pd.DataFrame(proj_rows).sort_values(["Topic ID", "Project Acronym"]).reset_index(drop=True)
print(f"Project rows: {len(projects_df)}")

TOPIC_WIDTHS = {
    "Framework": 18, "Year": 8, "Cluster/Pillar": 14, "Destination": 18, "Destination Name": 36,
    "Call ID": 32, "Call Title": 40, "Topic ID": 40, "Topic Title": 60, "Topic Type": 12,
    "TRL": 10, "EU Contribution per Project (EUR M)": 16, "Indicative Budget (EUR M)": 14,
    "Expected Number of Projects": 12, "Opening Date": 13, "Deadline": 13, "Stage": 14,
    "Keywords": 24, "EU Portal Link": 50, "Work Programme PDF": 24, "Notes": 40,
    "Call Status": 14, "Award Status": 14, "Funded Projects (signed)": 12,
    "Signed EU Contribution (EUR M)": 14,
}
PROJ_WIDTHS = {
    "Topic ID": 36, "Project Acronym": 18, "Project Full Name": 50, "Grant Agreement ID": 16,
    "Coordinator": 36, "Coordinator Country": 14, "Number of Participants": 12,
    "Participant Countries": 30, "Start Date": 12, "End Date": 12, "EU Contribution (EUR)": 16,
    "Total Cost (EUR)": 16, "Status": 12, "Project Website": 30, "CORDIS Link": 40,
    "EU Portal Link": 50, "Objective": 80, "Notes": 20,
}

def write_sheet(wb, name, df, widths, header_rgb):
    ws = wb.create_sheet(name)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=header_rgb)
    for c_idx, h in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = widths.get(h, 16)
    for r_idx, rec in enumerate(df.to_dict("records"), start=2):
        for c_idx, h in enumerate(df.columns, start=1):
            v = rec.get(h)
            if v is not None and not isinstance(v, (list, dict)) and pd.isna(v):
                v = None
            ws.cell(row=r_idx, column=c_idx, value=v)
    ws.freeze_panes = "A2"

wb = Workbook()
wb.remove(wb.active)
write_sheet(wb, "Topics", topics_df, TOPIC_WIDTHS, "2C5D99")
write_sheet(wb, "Projects", projects_df, PROJ_WIDTHS, "2F6B3D")
wb.save(OUT)
print(f"\nSaved: {OUT}")
print(f"  Topics sheet: {len(topics_df)} rows x {len(topics_df.columns)} cols")
print(f"  Projects sheet: {len(projects_df)} rows x {len(projects_df.columns)} cols")
