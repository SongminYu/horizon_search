"""Append/refresh an "Organisations" sheet: the Participants rows aggregated to
one row per unique organisation, with totals across all its projects in this
workbook's scope.

Usage: python3 pipeline/scripts/add_organisations_sheet.py <projects_pkl_stem> <workbook>
  e.g. python3 pipeline/scripts/add_organisations_sheet.py cl5_projects CL5_Topics_and_Projects.xlsx

Columns: organisation identity + #projects, #times coordinator, summed EC
contribution. Sorted by total EC contribution (descending), so the biggest
beneficiaries are on top. Only the Organisations sheet is touched.
"""
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
PARSED = ROOT / "parsed"

pkl_stem, wb_name = sys.argv[1], sys.argv[2]
RAW_SUBDIR = sys.argv[3] if len(sys.argv) > 3 else "cordis_he"
proj = pd.read_pickle(PARSED / f"{pkl_stem}.pkl")
XLSX = ROOT.parent / "data" / wb_name
pids = set(proj["id"])

org = pd.read_excel(ROOT / "raw" / RAW_SUBDIR / "organization.xlsx")
org = org[org["projectID"].isin(pids)].copy()

def to_float(x):
    if pd.isna(x):
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None

for col in ["ecContribution", "netEcContribution"]:
    org[col] = org[col].apply(to_float)

# organisationID is the stable key; name/country taken from the most recent row
org = org.sort_values("contentUpdateDate")
g = org.groupby("organisationID").agg(
    **{
        "Organisation Name": ("name", "last"),
        "Short Name": ("shortName", "last"),
        "Country": ("country", "last"),
        "City": ("city", "last"),
        "Activity Type": ("activityType", "last"),
        "SME": ("SME", "last"),
        "Projects": ("projectID", "nunique"),
        "Times Coordinator": ("role", lambda s: (s == "coordinator").sum()),
        "Total EC Contribution (EUR)": ("ecContribution", "sum"),
        "Total Net EC Contribution (EUR)": ("netEcContribution", "sum"),
    }
).reset_index().rename(columns={"organisationID": "Organisation ID"})

g = g.sort_values("Total EC Contribution (EUR)", ascending=False).reset_index(drop=True)
g.insert(0, "Rank", range(1, len(g) + 1))
print(f"Organisations: {len(g)} unique across {len(pids)} projects")
print(f"Total EC: EUR {g['Total EC Contribution (EUR)'].sum()/1e9:.2f}B")
print("Top 5:")
print(g.head(5)[["Organisation Name", "Country", "Projects", "Total EC Contribution (EUR)"]].to_string(index=False))

WIDTHS = {
    "Rank": 7, "Organisation ID": 14, "Organisation Name": 55, "Short Name": 16,
    "Country": 9, "City": 18, "Activity Type": 10, "SME": 8, "Projects": 10,
    "Times Coordinator": 12, "Total EC Contribution (EUR)": 17, "Total Net EC Contribution (EUR)": 17,
}

wb = load_workbook(XLSX)
if "Organisations" in wb.sheetnames:
    del wb["Organisations"]
ws = wb.create_sheet("Organisations")
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="6C3A85")
for c_idx, h in enumerate(g.columns, start=1):
    cell = ws.cell(row=1, column=c_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = WIDTHS.get(h, 14)
for r_idx, rec in enumerate(g.to_dict("records"), start=2):
    for c_idx, h in enumerate(g.columns, start=1):
        v = rec.get(h)
        if v is not None and not isinstance(v, (list, dict)) and pd.isna(v):
            v = None
        ws.cell(row=r_idx, column=c_idx, value=v)
ws.freeze_panes = "A2"
wb.save(XLSX)
print(f"Saved {XLSX} (Organisations sheet: {len(g)} rows x {len(g.columns)} cols)")
